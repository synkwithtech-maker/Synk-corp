"""
Multi-Source Document & URL Ingestion Service for Synk Corp Knowledge Base.
Supports:
1. PDF Documents (.pdf) with page-level metadata
2. Microsoft Word Documents (.docx)
3. Plain Text & Markdown (.txt, .md)
4. Tabular CSV Data (.csv)
5. Structured JSON Records (.json)
6. Web URLs (HTML web scraping and text extraction)
Indexes vectors into persistent Qdrant storage with batch chunking.
"""

import argparse
import csv
import json
import logging
import sys
import time
from pathlib import Path
from typing import List, Optional
from urllib.parse import urlparse

import docx
import pypdf
import requests
from bs4 import BeautifulSoup
from langchain_community.document_loaders import TextLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_qdrant import QdrantVectorStore

from app.core.config import get_settings

logger = logging.getLogger("synkcorp.ingest")


def load_text_or_markdown(file_path: Path) -> List[Document]:
    """Loads plain text or markdown files."""
    try:
        loader = TextLoader(str(file_path), encoding="utf-8")
        docs = loader.load()
        for doc in docs:
            doc.metadata["source"] = file_path.name
            doc.metadata["file_path"] = str(file_path)
            doc.metadata["file_type"] = file_path.suffix.lower()
        return docs
    except Exception as e:
        logger.error("Failed to load text file '%s': %s", file_path.name, e)
        return []


def load_pdf(file_path: Path) -> List[Document]:
    """
    Extracts text page-by-page from a PDF document using pypdf.
    Attaches page numbers to document metadata for precise citations.
    """
    docs: List[Document] = []
    try:
        reader = pypdf.PdfReader(str(file_path))
        total_pages = len(reader.pages)
        logger.info("Reading PDF '%s' (%d pages)", file_path.name, total_pages)

        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text and page_text.strip():
                docs.append(
                    Document(
                        page_content=page_text.strip(),
                        metadata={
                            "source": file_path.name,
                            "file_path": str(file_path),
                            "file_type": ".pdf",
                            "page": i + 1,
                            "total_pages": total_pages,
                        },
                    )
                )
        logger.info("Extracted %d valid page(s) from PDF '%s'", len(docs), file_path.name)
    except Exception as e:
        logger.error("Failed to extract PDF '%s': %s", file_path.name, e)
    return docs


def load_docx(file_path: Path) -> List[Document]:
    """Extracts text paragraphs from a Microsoft Word document (.docx)."""
    try:
        doc = docx.Document(str(file_path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        if not paragraphs:
            return []

        full_text = "\n\n".join(paragraphs)
        return [
            Document(
                page_content=full_text,
                metadata={
                    "source": file_path.name,
                    "file_path": str(file_path),
                    "file_type": ".docx",
                    "paragraph_count": len(paragraphs),
                },
            )
        ]
    except Exception as e:
        logger.error("Failed to extract DOCX '%s': %s", file_path.name, e)
        return []


def load_csv(file_path: Path) -> List[Document]:
    """Extracts structured rows from a CSV file into contextual documents."""
    docs: List[Document] = []
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.DictReader(f)
            for i, row in enumerate(reader):
                line = " | ".join(f"{k}: {v}" for k, v in row.items() if v)
                if line.strip():
                    docs.append(
                        Document(
                            page_content=line.strip(),
                            metadata={
                                "source": file_path.name,
                                "file_path": str(file_path),
                                "file_type": ".csv",
                                "row": i + 1,
                            },
                        )
                    )
        logger.info("Extracted %d rows from CSV '%s'", len(docs), file_path.name)
    except Exception as e:
        logger.error("Failed to parse CSV '%s': %s", file_path.name, e)
    return docs


def load_json(file_path: Path) -> List[Document]:
    """Extracts JSON records or nested objects into contextual documents."""
    docs: List[Document] = []
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            data = json.load(f)

        if isinstance(data, list):
            for i, item in enumerate(data):
                content = json.dumps(item, indent=2, ensure_ascii=False) if isinstance(item, dict) else str(item)
                docs.append(
                    Document(
                        page_content=content,
                        metadata={
                            "source": file_path.name,
                            "file_path": str(file_path),
                            "file_type": ".json",
                            "record_index": i + 1,
                        },
                    )
                )
        elif isinstance(data, dict):
            docs.append(
                Document(
                    page_content=json.dumps(data, indent=2, ensure_ascii=False),
                    metadata={
                        "source": file_path.name,
                        "file_path": str(file_path),
                        "file_type": ".json",
                    },
                )
            )
        logger.info("Extracted %d record(s) from JSON '%s'", len(docs), file_path.name)
    except Exception as e:
        logger.error("Failed to parse JSON '%s': %s", file_path.name, e)
    return docs


def load_document_file(file_path: Path) -> List[Document]:
    """Auto-detects document format and invokes the appropriate loader."""
    ext = file_path.suffix.lower()
    if ext in [".txt", ".md"]:
        return load_text_or_markdown(file_path)
    elif ext == ".pdf":
        return load_pdf(file_path)
    elif ext == ".docx":
        return load_docx(file_path)
    elif ext == ".csv":
        return load_csv(file_path)
    elif ext == ".json":
        return load_json(file_path)
    else:
        logger.warning("Unsupported file extension '%s' for '%s'", ext, file_path.name)
        return []


def load_url(url: str, timeout: int = 15) -> List[Document]:
    """
    Fetches HTML from a webpage, removes navigation/boilerplate,
    and extracts clean body text.
    """
    parsed = urlparse(url)
    if not parsed.scheme or not parsed.netloc:
        raise ValueError(f"Invalid URL format: '{url}'")

    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) SynkCorpIngestionBot/2.0",
        "Accept": "text/html,application/xhtml+xml",
    }

    logger.info("Fetching webpage: %s", url)
    response = requests.get(url, headers=headers, timeout=timeout)
    response.raise_for_status()

    soup = BeautifulSoup(response.text, "html.parser")

    for element in soup(["script", "style", "nav", "footer", "header", "aside", "noscript"]):
        element.decompose()

    page_title = soup.title.string.strip() if soup.title and soup.title.string else parsed.netloc
    clean_text = " ".join(soup.stripped_strings)

    if not clean_text:
        raise ValueError(f"No readable text could be extracted from {url}")

    logger.info("Extracted %d characters from URL '%s' (Title: '%s')", len(clean_text), url, page_title)

    return [
        Document(
            page_content=clean_text,
            metadata={
                "source": url,
                "title": page_title,
                "domain": parsed.netloc,
                "file_type": "url",
            },
        )
    ]


def load_documents_from_directory(doc_dir: Path) -> List[Document]:
    """Loads all supported files from a directory."""
    if not doc_dir.exists():
        raise FileNotFoundError(f"Documents directory does not exist: {doc_dir.resolve()}")

    supported_extensions = [".txt", ".md", ".pdf", ".docx", ".csv", ".json"]
    loaded_docs: List[Document] = []

    for file_path in doc_dir.iterdir():
        if file_path.is_file() and file_path.suffix.lower() in supported_extensions:
            docs = load_document_file(file_path)
            loaded_docs.extend(docs)

    return loaded_docs


def index_documents(documents: List[Document], force_recreate: bool = False) -> int:
    """Chunks documents, generates dense embeddings, and stores them in Qdrant."""
    if not documents:
        logger.warning("No documents provided to index.")
        return 0

    settings = get_settings()

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=700,
        chunk_overlap=100,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(documents)
    logger.info("Split %d document(s) into %d chunk(s)", len(documents), len(chunks))

    embeddings = HuggingFaceEmbeddings(
        model_name=settings.EMBEDDING_MODEL_NAME,
        model_kwargs={"device": "cpu"},
        encode_kwargs={"normalize_embeddings": True},
    )

    if force_recreate:
        logger.info("Rebuilding collection '%s' at '%s'", settings.QDRANT_COLLECTION_NAME, settings.QDRANT_PATH)
        QdrantVectorStore.from_documents(
            documents=chunks,
            embedding=embeddings,
            path=settings.QDRANT_PATH,
            collection_name=settings.QDRANT_COLLECTION_NAME,
            force_recreate=True,
        )
    else:
        logger.info("Appending chunks to existing collection '%s'", settings.QDRANT_COLLECTION_NAME)
        try:
            vs = QdrantVectorStore.from_existing_collection(
                embedding=embeddings,
                collection_name=settings.QDRANT_COLLECTION_NAME,
                path=settings.QDRANT_PATH,
            )
            vs.add_documents(chunks)
        except Exception:
            QdrantVectorStore.from_documents(
                documents=chunks,
                embedding=embeddings,
                path=settings.QDRANT_PATH,
                collection_name=settings.QDRANT_COLLECTION_NAME,
            )

    return len(chunks)


def ingest_knowledge_base() -> int:
    """Rebuilds the entire Qdrant vector database from the documents/ directory."""
    settings = get_settings()
    docs_path = settings.DOCUMENTS_DIR
    logger.info("Starting complete knowledge base ingestion from '%s'", docs_path.resolve())
    start_time = time.perf_counter()

    raw_documents = load_documents_from_directory(docs_path)
    if not raw_documents:
        raise ValueError(f"No valid documents found in {docs_path.resolve()}")

    chunk_count = index_documents(raw_documents, force_recreate=True)
    duration = round(time.perf_counter() - start_time, 2)
    logger.info("Knowledge base rebuild complete: %d chunks indexed in %s seconds.", chunk_count, duration)
    return chunk_count
